import csv
import io
import re
import uuid
from typing import Any, Optional

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.crm import Lead, LeadSource, LeadStage
from app.models.customers import ActivityType, Customer, CustomerActivity
from app.models.users import User
from app.schemas.customer import CustomerCreate
from app.services.ai_service import AIService

CUSTOMER_FIELDS = [
    "full_name_ar",
    "full_name_en",
    "mobile",
    "whatsapp",
    "email",
    "national_id",
    "passport_number",
    "address",
    "city",
    "country",
    "occupation",
    "employer",
    "income_range",
    "notes",
    "source",
    "create_lead",
]

FIELD_ALIASES: dict[str, list[str]] = {
    "full_name_en": [
        "name", "full name", "fullname", "client name", "customer name", "english name",
        "name en", "contact name", "client", "customer", "lead name", "prospect",
        "first name", "last name", "buyer", "owner name",
    ],
    "full_name_ar": [
        "الاسم", "الاسم بالعربي", "الاسم بالعربية", "اسم العميل", "العميل",
        "name ar", "arabic name", "arabic", "الاسم عربي",
    ],
    "mobile": [
        "mobile", "phone", "tel", "telephone", "cell", "cellphone", "contact number",
        "phone number", "mobile number", "رقم الموبايل", "الموبايل", "الهاتف", "رقم الهاتف",
        "رقم الجوال", "جوال", "موبايل", "phone no", "mobile no", "contact phone",
    ],
    "whatsapp": [
        "whatsapp", "whats app", "wa", "واتساب", "واتس", "whatsapp number",
    ],
    "email": [
        "email", "e-mail", "mail", "email address", "البريد", "البريد الالكتروني",
        "الايميل", "ايميل",
    ],
    "national_id": [
        "national id", "nationalid", "id number", "id no", "nid", "رقم قومي",
        "الرقم القومي", "بطاقة", "national number", "identity",
    ],
    "passport_number": [
        "passport", "passport number", "passport no", "جواز", "رقم الجواز",
    ],
    "address": [
        "address", "street", "location", "العنوان", "عنوان", "full address",
    ],
    "city": [
        "city", "town", "governorate", "محافظة", "المدينة", "مدينة", "area", "منطقة",
    ],
    "country": [
        "country", "nation", "الدولة", "بلد",
    ],
    "occupation": [
        "occupation", "job", "profession", "work", "الوظيفة", "مهنة", "job title",
    ],
    "employer": [
        "employer", "company", "workplace", "جهة العمل", "شركة", "employer name",
    ],
    "income_range": [
        "income", "salary", "income range", "budget", "الدخل", "الراتب", "ميزانية",
    ],
    "notes": [
        "notes", "note", "comments", "comment", "remarks", "description", "ملاحظات",
        "ملاحظة", "تعليق",
    ],
    "source": [
        "source", "lead source", "channel", "origin", "المصدر", "مصدر",
    ],
    "create_lead": [
        "create lead", "lead", "create_lead",
    ],
}

# In-memory preview sessions (single-worker deploy; confirm re-sends row data as fallback)
_preview_sessions: dict[str, dict[str, Any]] = {}


def _normalize_header(header: str) -> str:
    h = header.strip().lower()
    h = re.sub(r"[^\w\s\u0600-\u06FF]", " ", h)
    return re.sub(r"\s+", " ", h).strip()


def _has_arabic(text: str) -> bool:
    return bool(re.search(r"[\u0600-\u06FF]", text))


def _normalize_mobile(value: str) -> str:
    if not value:
        return ""
    raw = str(value).strip()
    if re.match(r"^\d+\.0$", raw):
        raw = raw[:-2]
    digits = re.sub(r"\D", "", raw)
    if not digits:
        return raw
    if digits.startswith("20") and len(digits) >= 12:
        return "0" + digits[2:13]
    if digits.startswith("2") and len(digits) == 11:
        return "0" + digits[1:]
    if len(digits) == 10 and digits.startswith("1"):
        return "0" + digits
    if len(digits) >= 11 and digits.startswith("0"):
        return digits[:11]
    return digits if len(digits) >= 8 else raw


def _score_header_match(norm_header: str, alias: str) -> float:
    if norm_header == alias:
        return 1.0
    if alias in norm_header or norm_header in alias:
        return 0.85
    h_tokens = set(norm_header.split())
    a_tokens = set(alias.split())
    if h_tokens & a_tokens:
        return 0.7
    return 0.0


def detect_column_mapping(headers: list[str], override: Optional[dict[str, str]] = None) -> dict[str, Any]:
    """Map file column headers to customer fields using aliases + heuristics."""
    if override:
        return {
            "mapping": override,
            "suggestions": [
                {"file_column": k, "mapped_to": v, "confidence": "manual"}
                for k, v in override.items()
            ],
            "unmapped": [h for h in headers if h not in override],
        }

    mapping: dict[str, str] = {}
    used_fields: set[str] = set()
    suggestions: list[dict[str, str]] = []

    for header in headers:
        norm = _normalize_header(header)
        if not norm:
            continue
        best_field = ""
        best_score = 0.0
        for field, aliases in FIELD_ALIASES.items():
            if field in used_fields:
                continue
            for alias in aliases:
                score = _score_header_match(norm, alias)
                if score > best_score:
                    best_score = score
                    best_field = field
        if best_field and best_score >= 0.7:
            mapping[header] = best_field
            used_fields.add(best_field)
            confidence = "high" if best_score >= 0.95 else "medium"
            suggestions.append({"file_column": header, "mapped_to": best_field, "confidence": confidence})

    unmapped = [h for h in headers if h not in mapping]

    # Single generic "name" → fill both AR and EN if only one name column
    name_headers = [h for h, f in mapping.items() if f == "full_name_en"]
    if name_headers and "full_name_ar" not in used_fields:
        mapping[name_headers[0] + "::__ar_mirror"] = "full_name_ar"

    return {"mapping": mapping, "suggestions": suggestions, "unmapped": unmapped}


def parse_upload(content: bytes, filename: str) -> tuple[list[str], list[dict[str, str]]]:
    """Parse CSV or XLSX into headers + row dicts."""
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xls")):
        return _parse_xlsx(content)
    return _parse_csv(content)


def _parse_csv(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    for encoding in ("utf-8-sig", "utf-8", "cp1256", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = content.decode("utf-8", errors="replace")

    sample = text[:4096]
    delimiter = ","
    if sample.count(";") > sample.count(","):
        delimiter = ";"
    elif sample.count("\t") > sample.count(","):
        delimiter = "\t"

    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    headers = [h for h in (reader.fieldnames or []) if h and h.strip()]
    rows = []
    for row in reader:
        cleaned = {k.strip(): str(v).strip() if v is not None else "" for k, v in row.items() if k}
        if any(v for v in cleaned.values()):
            rows.append(cleaned)
    return headers, rows


def _parse_xlsx(content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return [], []
    headers = [str(h).strip() if h is not None else f"column_{i}" for i, h in enumerate(header_row)]
    headers = [h for h in headers if h]
    rows: list[dict[str, str]] = []
    for row in rows_iter:
        values = list(row)
        if not any(v is not None and str(v).strip() for v in values):
            continue
        item: dict[str, str] = {}
        for i, h in enumerate(headers):
            if i < len(values) and values[i] is not None:
                val = values[i]
                if isinstance(val, float) and val == int(val):
                    val = int(val)
                item[h] = str(val).strip()
            else:
                item[h] = ""
        rows.append(item)
    wb.close()
    return headers, rows


def _apply_mapping(row: dict[str, str], mapping: dict[str, str]) -> dict[str, str]:
    out: dict[str, str] = {}
    for file_col, field in mapping.items():
        if file_col.endswith("::__ar_mirror"):
            src = file_col.replace("::__ar_mirror", "")
            val = row.get(src, "").strip()
            if val:
                out["full_name_ar"] = val
            continue
        val = row.get(file_col, "").strip()
        if val:
            out[field] = val

    # Smart name splitting
    name_val = out.get("full_name_en") or out.get("full_name_ar") or ""
    if name_val:
        if _has_arabic(name_val):
            if not out.get("full_name_ar"):
                out["full_name_ar"] = name_val
            if not out.get("full_name_en"):
                out["full_name_en"] = name_val
        else:
            if not out.get("full_name_en"):
                out["full_name_en"] = name_val
            if not out.get("full_name_ar"):
                out["full_name_ar"] = name_val

    if out.get("mobile"):
        out["mobile"] = _normalize_mobile(out["mobile"])
    if out.get("whatsapp"):
        out["whatsapp"] = _normalize_mobile(out["whatsapp"])
    elif out.get("mobile"):
        out["whatsapp"] = out["mobile"]

    if not out.get("country"):
        out["country"] = "Egypt"

    return out


class CustomerImportService:
    MAX_ROWS = 2000

    @staticmethod
    def preview_file(
        content: bytes,
        filename: str,
        column_mapping: Optional[dict[str, str]] = None,
    ) -> dict[str, Any]:
        headers, raw_rows = parse_upload(content, filename)
        if not headers:
            return {"error": "No columns found in file"}
        if len(raw_rows) > CustomerImportService.MAX_ROWS:
            raw_rows = raw_rows[: CustomerImportService.MAX_ROWS]

        detection = detect_column_mapping(headers, column_mapping)
        mapping = detection["mapping"]
        preview_rows: list[dict[str, Any]] = []

        for idx, raw in enumerate(raw_rows, start=2):
            mapped = _apply_mapping(raw, mapping)
            errors: list[str] = []
            status = "valid"

            if not mapped.get("full_name_en") and not mapped.get("full_name_ar"):
                errors.append("Name is required")
            if not mapped.get("mobile"):
                errors.append("Mobile is required")
            elif len(mapped.get("mobile", "")) < 8:
                errors.append("Mobile number too short")

            try:
                if not errors:
                    CustomerCreate(
                        full_name_ar=mapped.get("full_name_ar") or mapped.get("full_name_en", ""),
                        full_name_en=mapped.get("full_name_en") or mapped.get("full_name_ar", ""),
                        mobile=mapped.get("mobile", ""),
                        whatsapp=mapped.get("whatsapp"),
                        email=mapped.get("email") or None,
                        national_id=mapped.get("national_id"),
                        passport_number=mapped.get("passport_number"),
                        address=mapped.get("address"),
                        city=mapped.get("city"),
                        country=mapped.get("country", "Egypt"),
                        occupation=mapped.get("occupation"),
                        employer=mapped.get("employer"),
                        income_range=mapped.get("income_range"),
                        notes=mapped.get("notes"),
                    )
            except ValidationError as e:
                for err in e.errors():
                    errors.append(err.get("msg", "Validation error"))

            if errors:
                status = "error"

            preview_rows.append({
                "row_number": idx,
                "status": status,
                "errors": errors,
                "data": mapped,
                "raw": raw,
            })

        session_id = str(uuid.uuid4())
        _preview_sessions[session_id] = {
            "mapping": mapping,
            "rows": preview_rows,
            "filename": filename,
        }

        valid = sum(1 for r in preview_rows if r["status"] == "valid")
        errors_count = sum(1 for r in preview_rows if r["status"] == "error")

        return {
            "session_id": session_id,
            "filename": filename,
            "headers": headers,
            "available_fields": CUSTOMER_FIELDS,
            "column_mapping": mapping,
            "suggestions": detection["suggestions"],
            "unmapped_columns": detection["unmapped"],
            "rows": preview_rows,
            "summary": {
                "total": len(preview_rows),
                "valid": valid,
                "errors": errors_count,
            },
        }

    @staticmethod
    async def check_duplicates(
        db: AsyncSession,
        company_id: int,
        rows: list[dict[str, Any]],
    ) -> list[dict[str, Any]]:
        mobiles = [r["data"]["mobile"] for r in rows if r.get("data", {}).get("mobile")]
        national_ids = [r["data"]["national_id"] for r in rows if r.get("data", {}).get("national_id")]
        existing_by_mobile: dict[str, Customer] = {}
        existing_by_nid: dict[str, Customer] = {}

        if mobiles or national_ids:
            conditions = []
            if mobiles:
                conditions.append(Customer.mobile.in_(mobiles))
            if national_ids:
                conditions.append(Customer.national_id.in_(national_ids))
            result = await db.execute(
                select(Customer).where(
                    Customer.company_id == company_id,
                    Customer.is_deleted == False,
                    or_(*conditions),
                )
            )
            for c in result.scalars().all():
                existing_by_mobile[c.mobile] = c
                if c.national_id:
                    existing_by_nid[c.national_id] = c

        updated = []
        for row in rows:
            data = row.get("data", {})
            mobile = data.get("mobile", "")
            nid = data.get("national_id", "")
            existing = existing_by_mobile.get(mobile) or (existing_by_nid.get(nid) if nid else None)
            r = dict(row)
            if existing and r["status"] == "valid":
                r["status"] = "duplicate"
                r["existing_id"] = existing.id
                r["existing_name"] = existing.full_name_en
            updated.append(r)
        return updated

    @staticmethod
    async def confirm_import(
        db: AsyncSession,
        company_id: int,
        user: User,
        rows: list[dict[str, Any]],
        duplicate_strategy: str = "skip",
        create_leads: bool = False,
        default_source: str = "walk_in",
    ) -> dict[str, Any]:
        created = 0
        updated = 0
        skipped = 0
        failed = 0
        results: list[dict[str, Any]] = []

        for row in rows:
            row_num = row.get("row_number", 0)
            if row.get("status") == "error":
                failed += 1
                results.append({"row_number": row_num, "status": "failed", "reason": "; ".join(row.get("errors", []))})
                continue

            data = row.get("data", {})
            existing_id = row.get("existing_id")

            if row.get("status") == "duplicate":
                if duplicate_strategy == "skip":
                    skipped += 1
                    results.append({"row_number": row_num, "status": "skipped", "reason": "Duplicate mobile"})
                    continue
                if duplicate_strategy == "update" and existing_id:
                    cust = (await db.execute(select(Customer).where(Customer.id == existing_id))).scalar_one_or_none()
                    if cust:
                        for field in CUSTOMER_FIELDS:
                            if field in ("source", "create_lead"):
                                continue
                            val = data.get(field)
                            if val:
                                setattr(cust, field, val)
                        updated += 1
                        results.append({"row_number": row_num, "status": "updated", "customer_id": cust.id})
                        continue

            try:
                payload = CustomerCreate(
                    full_name_ar=data.get("full_name_ar") or data.get("full_name_en", ""),
                    full_name_en=data.get("full_name_en") or data.get("full_name_ar", ""),
                    mobile=data["mobile"],
                    whatsapp=data.get("whatsapp"),
                    email=data.get("email") or None,
                    national_id=data.get("national_id"),
                    passport_number=data.get("passport_number"),
                    address=data.get("address"),
                    city=data.get("city"),
                    country=data.get("country", "Egypt"),
                    occupation=data.get("occupation"),
                    employer=data.get("employer"),
                    income_range=data.get("income_range"),
                    notes=data.get("notes"),
                    assigned_to_id=user.id,
                )
            except ValidationError as e:
                failed += 1
                results.append({"row_number": row_num, "status": "failed", "reason": str(e)})
                continue

            customer = Customer(company_id=company_id, **payload.model_dump())
            db.add(customer)
            await db.flush()

            db.add(
                CustomerActivity(
                    customer_id=customer.id,
                    activity_type=ActivityType.NOTE,
                    title="Imported from bulk upload",
                    description=f"Row {row_num} imported via smart bulk upload",
                    created_by_id=user.id,
                )
            )

            should_lead = create_leads or str(data.get("create_lead", "")).lower() in ("true", "1", "yes")
            if should_lead:
                try:
                    source = LeadSource(data.get("source") or default_source)
                except ValueError:
                    source = LeadSource.WALK_IN
                lead = Lead(
                    company_id=company_id,
                    customer_id=customer.id,
                    full_name=customer.full_name_en,
                    mobile=customer.mobile,
                    email=customer.email,
                    source=source,
                    stage=LeadStage.NEW,
                    assigned_to_id=user.id,
                    notes=customer.notes,
                )
                score, ai_score, qual = AIService.score_lead(lead)
                lead.score = score
                lead.ai_score = ai_score
                lead.ai_qualification = qual
                db.add(lead)

            created += 1
            results.append({"row_number": row_num, "status": "created", "customer_id": customer.id})

        await db.flush()
        return {
            "created": created,
            "updated": updated,
            "skipped": skipped,
            "failed": failed,
            "results": results,
        }

    @staticmethod
    def get_template_rows() -> list[dict[str, str]]:
        return [
            {
                "full_name_ar": "أحمد محمد",
                "full_name_en": "Ahmed Mohamed",
                "mobile": "01012345678",
                "whatsapp": "01012345678",
                "email": "ahmed@example.com",
                "national_id": "29001011234567",
                "city": "Cairo",
                "country": "Egypt",
                "occupation": "Engineer",
                "notes": "Interested in 3BR apartment",
                "source": "facebook",
            },
            {
                "full_name_ar": "سارة علي",
                "full_name_en": "Sara Ali",
                "mobile": "01098765432",
                "email": "sara@example.com",
                "city": "Giza",
                "country": "Egypt",
                "source": "referral",
            },
        ]

    @staticmethod
    def get_session(session_id: str) -> Optional[dict[str, Any]]:
        return _preview_sessions.get(session_id)
